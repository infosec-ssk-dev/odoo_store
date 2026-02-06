import os
from odoo import http
from odoo.http import request
import io, base64, copy
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from datetime import datetime


class DdAExcelReportEbiosRM(http.Controller):

    def _load_template_workbook(self, study_id, template_id=None):
        """
        Load template workbook with priority:
        1. Specific template_id if provided
        2. Most recent custom template with use_this_template=True
        3. System default template
        4. Fallback to file system
        """

        # If specific template requested, use it
        if template_id:
            tmpl = request.env['digiit.ebios_rm.report.template'].browse(template_id)
            if tmpl.exists() and tmpl.template_file:
                data = base64.b64decode(tmpl.template_file)
                return load_workbook(io.BytesIO(data))

        # Priority 1: Look for custom templates that are enabled (use_this_template=True)
        tmpl = request.env['digiit.ebios_rm.report.template'].search([
            ('report_type', '=', 'dda_ebios'),
            ('is_default', '=', False),
            ('use_this_template', '=', True)
        ], limit=1, order='create_date desc')

        if tmpl and tmpl.template_file:
            data = base64.b64decode(tmpl.template_file)
            return load_workbook(io.BytesIO(data))

        # Priority 2: Look for system default template
        tmpl = request.env['digiit.ebios_rm.report.template'].search([
            ('report_type', '=', 'dda_ebios'),
            ('is_default', '=', True)
        ], limit=1)

        if tmpl and tmpl.template_file:
            data = base64.b64decode(tmpl.template_file)
            return load_workbook(io.BytesIO(data))

        # Fallback: Load from file system (legacy support)
        module_path = os.path.dirname(os.path.abspath(__file__))
        default_path = os.path.join(module_path, '..', 'static', 'templates', 'DdA-SMSI-v0.1.xlsx')

        if not os.path.isfile(default_path):
            raise FileNotFoundError(
                "Aucun template trouvé. Veuillez configurer un template dans les paramètres."
            )

        return load_workbook(default_path)

    def _safe_cell(self, sheet, row, col):

        cell = sheet.cell(row=row, column=col)

        if not isinstance(cell, MergedCell):
            return cell

        # Identify the merged block this cell belongs to
        for merged_range in sheet.merged_cells.ranges:
            if merged_range.min_row <= row <= merged_range.max_row and \
               merged_range.min_col <= col <= merged_range.max_col:

                # Return top-left cell of the merge
                return sheet.cell(row=merged_range.min_row, column=merged_range.min_col)

        return cell

    def _get_sheet(self, workbook, expected_names):

        expected = [e.lower() for e in expected_names]

        for sheet in workbook.worksheets:
            name = sheet.title.lower()
            if all(token in name for token in expected):
                return sheet

        return None

    @http.route('/ebios_rm/dda/excel/report/<int:study_id>', type='http', auth='user')
    def download_dda_ebios_rm_excel_report(self, study_id, template_id=None, **kw):

        try:
            workbook = self._load_template_workbook(study_id, template_id)
        except Exception as e:
            return request.make_response(
                f"Impossible de charger le modèle : {e}",
                headers=[('Content-Type', 'text/plain')]
            )

        cover_sheet = self._get_sheet(workbook, ["f1"])
        data_sheet = self._get_sheet(workbook, ["soa", "french", "2022"])

        if not data_sheet:
            return request.make_response(
                "Erreur: la feuille DdA 'SoA French 2022' n’a pas été trouvée dans le modèle.",
                headers=[('Content-Type', 'text/plain')]
            )

        study = request.env['digiit.ebios_rm.study'].browse(study_id)
        paragraphs_ids = study.study_measure_ids

        if cover_sheet:
            self._fill_cover_page(cover_sheet, study)
        self._fill_data_sheet(data_sheet, paragraphs_ids, study_id)

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        file_name = f"DdA_Excel_Report_{study.name or 'EbiosRM'}_{datetime.now():%Y%m%d}.xlsx"

        return request.make_response(
            output.getvalue(),
            headers=[
                ('Content-Type',
                 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
                ('Content-Disposition', f'attachment; filename="{file_name}"')
            ]
        )

    def _merge_chapter_cells(self, sheet, start_row, end_row):
        """
        Merge cells in column A (Chapter) for consecutive rows with same chapter.
        """
        from openpyxl.styles import Alignment

        if start_row >= end_row:
            return

        current_chapter = None
        merge_start_row = start_row

        for row in range(start_row, end_row + 2):  # +2 to process the last group
            if row <= end_row:
                cell_value = self._safe_cell(sheet, row, 1).value
            else:
                cell_value = None  # Trigger final merge

            if cell_value != current_chapter:
                # Merge previous group if it spans multiple rows
                if current_chapter is not None and row - 1 > merge_start_row:
                    sheet.merge_cells(
                        start_row=merge_start_row,
                        start_column=1,
                        end_row=row - 1,
                        end_column=1
                    )

                    merged_cell = sheet.cell(row=merge_start_row, column=1)
                    old_alignment = merged_cell.alignment
                    merged_cell.alignment = Alignment(
                        horizontal=old_alignment.horizontal,
                        vertical='top',
                        text_rotation=old_alignment.text_rotation,
                        wrap_text=old_alignment.wrap_text,
                        shrink_to_fit=old_alignment.shrink_to_fit,
                        indent=old_alignment.indent
                    )

                # Start new group
                current_chapter = cell_value
                merge_start_row = row

    def _fill_cover_page(self, sheet, study):
        today = datetime.now().strftime('%d/%m/%Y')
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{today}}" in cell.value:
                    cell.value = cell.value.replace("{{today}}", today)

    def _fill_data_sheet(self, sheet, paragraphs_ids, study_id):
        today = datetime.now().strftime('%d/%m/%Y')

        # Replace dates in header
        for row in sheet.iter_rows(min_row=1, max_row=3):
            for cell in row:
                if isinstance(cell.value, str) and "{{today}}" in cell.value:
                    cell.value = cell.value.replace("{{today}}", today)

        start_row = 4
        template_row = start_row
        row_num = start_row

        study = request.env['digiit.ebios_rm.study'].browse(study_id)

        for paragraph in paragraphs_ids:

            self._safe_cell(sheet, row_num,1).value = f"{paragraph.study_chapter_id.chapter_number} - {paragraph.study_chapter_id.chapter_title}"
            self._safe_cell(sheet, row_num, 2).value = paragraph.measure_number
            self._safe_cell(sheet, row_num, 3).value = paragraph.measure_title
            self._safe_cell(sheet, row_num, 5).value = paragraph.measure_contenu

            # Risk treatments
            risk_treatments = request.env['digiit.ebios_rm.risk.treatment'].search([
                ('equivalent_mesures_referentiel', 'in', paragraph.id),
                ('study_id', '=', study_id)
            ])

            # Applicability
            if risk_treatments:
                applicable = "Oui"
                em_bp = "X"
                rar = "X"
                risks = ", ".join(rt.identifier or "N/A" for rt in risk_treatments)
                justification = f"Présence des risques suivant: {risks}"
            elif study.is_applicable:
                applicable = "Oui"
                em_bp = "X"
                rar = ""
                justification = "Bonne pratique"
            else:
                applicable = "Non"
                em_bp = ""
                rar = ""
                justification = "Aucun risque associé"

            # Fill columns F–K
            self._safe_cell(sheet, row_num, 6).value = applicable
            self._safe_cell(sheet, row_num, 7).value = justification
            self._safe_cell(sheet, row_num, 8).value = ""
            self._safe_cell(sheet, row_num, 9).value = em_bp
            self._safe_cell(sheet, row_num, 10).value = rar
            self._safe_cell(sheet, row_num, 11).value = ""

            # Copy styling from template row
            for col in range(1, 11):
                src = sheet.cell(row=template_row, column=col)
                tgt = sheet.cell(row=row_num, column=col)
                tgt.border = copy.copy(src.border)
                tgt.alignment = copy.copy(src.alignment)
                tgt.font = copy.copy(src.font)
                tgt.fill = copy.copy(src.fill)

            row_num += 1

        self._merge_chapter_cells(sheet, start_row, row_num - 1)